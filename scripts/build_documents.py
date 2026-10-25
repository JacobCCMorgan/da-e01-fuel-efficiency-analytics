"""Build the Chinese analysis report and English summary from report_content.json.

Requires python-docx and the Noto Sans CJK TC font for Chinese layout. Numerical
findings come from the structured report, never from independently typed results.
The optional PDF export uses LibreOffice; inspect rendered pages before release.
"""
from __future__ import annotations
import argparse
import json
import re
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BLACK="000000"; NAVY="163247"; GRAY="D9D9D9"; PALE="F1F5F7"
ZH_TITLES=["效率必須連同觀測覆蓋解讀","從車輛週矩陣定位變化","用載重與任務解釋耗用差異",
           "拆解購油支出的價格與數量效果","以油箱平衡交叉核對 ECU 耗用","先檢查可比性再安排調查","從報告採用追蹤到可比的後續結果"]
NEXT_ACTIONS=[
    "當效率曲線改變時，先查看同週配對時間覆蓋，再按車輛及任務下鑽。短期改善若伴隨觀測缺口，應先查量測與歸屬，不直接當作駕駛或車況變好。",
    "先沿同一列比較同車的時間變化，再選取同任務與車型的其他車輛。灰格表示沒有可用 MPG，可能是沒有營運暴露；其語意必須由明細狀態確認。",
    "載重採任務時間加權，氣泡大小呈現 PTO 占比。此圖提供分組背景；真正的同儕比較仍要拆到任務粒度，不能把 mixed 當成單一同質任務。",
    "交易淨額先排除不適用的金額與數量口徑，再計算價量效果。市場背景按區域及觀測週配對；交易條件、稅務口徑及採購可得性仍可能不同，市場差額不是可實現節省。",
    "先核對盤點兩端與燃料移動是否屬同一油箱範圍，再查看 ECU 區間覆蓋。超出不確定範圍的差額應回查補傳、未校準讀值及移動憑證；它本身不能證明偷油或浪費。",
    "同儕必須來自相同租戶、車型、任務與週，排除當前車輛。資料充足才比較差異；沒有足夠同儕表示證據不足，不能視為正常或異常的已驗證結論。",
    "先確認報告是否被採用與措施是否發生，再檢查觀察是否成熟、是否失聯及任務是否可比。保留原始與任務組合標準化結果；標準化能處理已知任務占比，仍未控制所有混雜因素。"]
LABELS={"within_measurement_uncertainty":"落於量測不確定範圍","no_eligible_ecu_intervals":"沒有合格 ECU 區間",
        "uncalibrated_boundary":"邊界量測未校準","incomplete_movements":"燃料移動紀錄不完整",
        "investigate_coverage_or_balance":"需調查覆蓋或庫存平衡","comparable_descriptive_only":"可作描述性比較",
        "no_action_observed":"未觀察到措施","task_mix_changed":"任務組合改變","lost_to_followup":"失去後續追蹤",
        "not_mature":"觀察窗口尚未成熟","linehaul":"長途幹線","regional_delivery":"區域配送","mixed":"混合任務"}

def clean(text):
    return str(text).replace("–","-").replace("—","-").replace("−","-")

def cn(text):
    for key,value in LABELS.items():text=text.replace(key,value)
    return clean(text)

def font(run,name,size=None,bold=None,color=BLACK):
    run.font.name=name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),name)
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)

def prepare(lang):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.62);sec.bottom_margin=Inches(.63)
    sec.left_margin=sec.right_margin=Inches(.65)
    sec.footer_distance=Inches(.25)
    family="Noto Sans CJK TC" if lang=="ZH" else "Arial"
    for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3","Caption"]:
        st=doc.styles[name];st.font.name=family;st.font.color.rgb=RGBColor.from_string(BLACK)
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),family)
        st.paragraph_format.space_after=Pt(7)
    doc.styles["Normal"].font.size=Pt(11.5 if lang=="ZH" else 11)
    doc.styles["Normal"].paragraph_format.line_spacing=1.24
    doc.styles["Title"].font.size=Pt(25 if lang=="ZH" else 26)
    doc.styles["Title"].font.bold=True
    doc.styles["Title"].paragraph_format.space_after=Pt(12)
    doc.styles["Subtitle"].font.size=Pt(12)
    doc.styles["Subtitle"].font.italic=False
    for name,size in [("Heading 1",17),("Heading 2",12.5),("Heading 3",11.5)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(9)
        doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles["Caption"].font.size=Pt(9.3)
    doc.styles["Caption"].font.italic=False
    doc.styles["Caption"].font.bold=False
    doc.styles["Caption"].paragraph_format.line_spacing=1.12
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(f.add_run("DA-E01  |  v1.0  |  "),family,8.5)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");f._p.append(field)
    doc.core_properties.title="車隊燃油效率與價格差異分析" if lang=="ZH" else "Fleet fuel efficiency and price variance"
    doc.core_properties.subject="DA-E01 analysis"
    doc.core_properties.author=""
    doc.core_properties.last_modified_by=""
    for st in doc.styles:
        for border in st._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    return doc,family

def para(doc,text,style=None):
    p=doc.add_paragraph(clean(text),style=style)
    p.paragraph_format.widow_control=True
    return p

def title(doc,text,subtitle=None):
    para(doc,text,"Title")
    if subtitle:para(doc,subtitle,"Subtitle")

def heading(doc,text):return para(doc,text,"Heading 1")

def newpage(doc,text):
    p=heading(doc,text);p.paragraph_format.page_break_before=True
    return p

def table(doc,headers,rows,widths,family):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
    for i,value in enumerate(headers):t.rows[0].cells[i].text=clean(value)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):cells[i].text=clean(value)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();keep=OxmlElement("w:cantSplit");trpr.append(keep)
        if ri==0:
            rep=OxmlElement("w:tblHeader");trpr.append(rep)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();shd=OxmlElement("w:shd")
            shd.set(qn("w:fill"),NAVY if ri==0 else (PALE if ri%2==0 else "FFFFFF"));pr.append(shd)
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                el=OxmlElement("w:"+edge);el.set(qn("w:val"),"single");el.set(qn("w:sz"),"4");el.set(qn("w:color"),GRAY);borders.append(el)
            pr.append(borders);m=OxmlElement("w:tcMar")
            for key,value in [("top",75),("bottom",75),("left",90),("right",90)]:
                el=OxmlElement("w:"+key);el.set(qn("w:w"),str(value));el.set(qn("w:type"),"dxa");m.append(el)
            pr.append(m)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(headers)==3 and ci==2 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:font(r,family,10.3,ri==0,"FFFFFF" if ri==0 else BLACK)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def figure(doc,root,fig,width=7.2):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
    r=p.add_run();r.add_picture(str(root/fig["path"]),width=Inches(width))
    for node in r._r.xpath(".//wp:docPr"):
        node.set("descr",fig["title"])
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def statuses(fig):
    out=[]
    for text in fig["findings"]:
        m=re.match(r"([^：]+)：(\d+) (?:個庫存窗口|筆審閱案例)。",text)
        if m:out.append((m[1],int(m[2])))
    return out

def cases(fig):
    out=[]
    pattern=r"(T\d+-V\d+) 的 ([^：]+)：基準 ([\d.]+)、原始後續 ([\d.]+)、按基準任務組合標準化後續 ([\d.]+) L/100km；標準化差異 ([+-][\d.]+)，任務里程占比最大變動 ([\d.]+)"
    for text in fig["findings"]:
        m=re.search(pattern,text)
        if m:out.append(m.groups())
    return out

def zh_report(root,r):
    doc,family=prepare("ZH");s=r["summary"]
    title(doc,"車隊燃油效率與價格差異分析","DA E01 專案分析報告")
    para(doc,"分析期間為 2025 年 1 至 6 月。營運指標採修訂版，公開市場價格用於事後背景比較。")
    para(doc,f"本報告把 ECU 耗用、購油交易與油箱存量分開核對，再按車輛、任務及資料覆蓋判讀差異。整體效率為 {s['fleet_mpg']:.2f} MPG，但調查優先順序取決於證據是否可比；購油差額或前後變化本身不能證明浪費或節省。")
    table(doc,["主要數值","本次結果","判讀口徑"],[
        ["車隊與觀察規模",f"{s['tenant_count']} 租戶  {s['vehicle_count']} 車",f"{s['report_week_count']} 報告週，{s['vehicle_week_rows']} 筆車輛週"],
        ["可配對里程",f"{s['paired_distance_km']:,.1f} km","與耗用使用相同合格區間"],
        ["ECU 配對耗用",f"{s['fuel_used_l']:,.1f} L","不以購油量代替"],
        ["車隊燃油效率",f"{s['fleet_mpg']:.2f} MPG","總里程除以總耗用後換算"],
        ["配對時間覆蓋",f"{s['paired_time_coverage']:.2%}","有效配對秒數除以應觀察秒數"],
        ["可比購油淨額",f"US${s['comparable_purchase_amount_usd']:,.2f}","交易支出，不是當期耗用成本"],
    ],[1.55,1.8,3.85],family)
    para(doc,"主要分析發現","Heading 2")
    para(doc,f"{s['low_coverage_rows']} 筆車輛週需要優先查核覆蓋；{s['no_eligible_peer_rows']} 筆缺少合格同儕基準。符合任務同儕規則的強度查核有 {s['intensity_review_rows']} 筆，價格查核有 {s['price_review_rows']} 筆，類別可能重疊。")
    para(doc,f"報告使用流程包含 {s['review_case_count']} 筆審閱案例，其中 {s['comparable_followup_cases']} 筆可作描述性後續比較。按基準任務里程占比標準化後，仍需保留未控制的路線、載重與其他情境差異。")
    para(doc,"資料性質與結論邊界","Heading 2")
    para(doc,"車輛營運、交易與措施追蹤為明確標示的合成資料；EIA 價格為獨立的公開市場背景。本案展示可重現的分析流程與模擬情境，不代表真實客戶採用、實際節油或公司營運成果。")
    for i,f in enumerate(r["figures"]):
        newpage(doc,ZH_TITLES[i])
        para(doc,f["interpretation"].replace("灰格代表未知", "灰格表示沒有可用效率"))
        figure(doc,root,f)
        if i==4:
            table(doc,["庫存核對狀態","窗口數"],[(LABELS.get(k,k),v) for k,v in statuses(f)],[5.6,1.6],family)
        elif i==6:
            state="；".join(f"{LABELS.get(k,k)} {v} 筆" for k,v in statuses(f))
            para(doc,state+"。")
            cs=cases(f);assert len(cs)==s["comparable_followup_cases"]
            table(doc,["車輛","基準","原始後續","標準化後續","標準化差異"],
                  [[x[0],x[2],x[3],x[4],x[5]] for x in cs],[1.4,1.2,1.4,1.6,1.6],family)
            para(doc,"耗用強度單位均為 L/100km。標準化以基準期各任務的里程占比加權後續任務強度；圖中空心菱形為原始後續值。","Caption")
            para(doc,f"活躍審閱合計 {s['recorded_active_review_seconds']/60:.1f} 分鐘，記錄基準為 {s['recorded_baseline_review_seconds']/60:.1f} 分鐘，涵蓋 {s['timed_review_cases']} 筆案例。活躍工時不含等待；這些合成紀錄只驗證流程處理。")
        else:
            for finding in f["findings"]:para(doc,cn(finding))
        para(doc,"分析後的查核方向","Heading 2")
        para(doc,NEXT_ACTIONS[i])
        if i==0:para(doc,f"另有 {s['no_operating_exposure_rows']} 筆沒有預定營運時間，效率不適用；{s['unknown_coverage_rows']} 筆覆蓋未知。這兩種狀態與低覆蓋分開保存。")
        if i==5:para(doc,f"缺少合格同儕的 {s['no_eligible_peer_rows']} 筆車輛週保留在明細中。未觸發規則不等於已證實正常；規則的用途是安排下一步查核。")
        if i!=6:para(doc,"圖表與明細均可由 output_id、車輛及報告週回查；完整計算口徑見 docs/metrics.md。","Caption")
    newpage(doc,"重現結果與使用分析證據")
    para(doc,"讀者可以由來源檔案與建置入口重建資料庫，再從 SQL 計算、圖表和明細回查同一結果。修改規則時應重新執行分析與驗證，保留輸入版本、參數及新的結果。")
    steps=[
        ("確認問題與口徑","先分清購油支出、ECU 耗用與庫存平衡；明確定義同一報告版本、時間窗口與分母。公開價格與合成營運資料各自保留來源，事件時間不等於平台可取得時間。","docs/metrics.md  data/public/  config/  contracts/"),
        ("檢查資料準備與品質","從原始事件、標準化欄位到共同有效區間，保留拒收、修訂、缺值及核對依據。區分量測不完整、同儕不足及窗口未成熟，避免將未知解讀成營運結果。","src/fleet_fuel/  validation/"),
        ("執行分析與調整規則","先核對總量，再計算配對效率、價量差異、同任務基準與可比的後續窗口。修改門檻後重新檢查涵蓋率與例外，並保留新的參數與結果。","sql/analytics.sql  notebooks/01_fuel_analysis.ipynb"),
        ("閱讀圖表與回查明細","互動報表一次選一個版本，再篩選租戶、車輛、週與品質並下載明細。利用 output_id 回查來源與截止時間；中文圖表說明提供每張圖的判讀與限制。","outputs/dashboard.html  outputs/report_weekly_detail.csv\ndocs/chart_walkthrough_zh.md"),
    ]
    for h,body,path in steps:
        para(doc,h,"Heading 2");para(doc,body);para(doc,path,"Caption")
    para(doc,"執行入口","Heading 2")
    para(doc,"依 README 安裝已鎖定的執行環境後，在專案根目錄執行：")
    p=para(doc,"python -m fleet_fuel.cli run --profile full --operational snapshot")
    for run in p.runs:font(run,"DejaVu Sans Mono",9)
    path=root/"reports/DA-E01_Analysis_Report_ZH.docx";doc.save(path);return path

def en_report(root,r):
    doc,family=prepare("EN");s=r["summary"];fs=r["figures"]
    title(doc,"Fleet fuel efficiency and price variance","DA E01 analytical summary")
    para(doc,"January to June 2025 operations. Revised operational view with separately identified retrospective market context.")
    para(doc,f"The fleet achieves {s['fleet_mpg']:.2f} MPG over matched measurement intervals. The decision is where evidence supports further investigation: purchased fuel, ECU consumption and tank stock are reconciled separately before efficiency or price differences are interpreted.")
    table(doc,["Measure","Result","Definition"],[
        ["Scope",f"{s['tenant_count']} tenants / {s['vehicle_count']} vehicles",f"{s['vehicle_week_rows']} vehicle-week rows"],
        ["Paired distance",f"{s['paired_distance_km']:,.1f} km","Accepted common intervals"],
        ["Paired ECU fuel",f"{s['fuel_used_l']:,.1f} L","Consumption, not purchases"],
        ["Fleet efficiency",f"{s['fleet_mpg']:.2f} MPG","Ratio of summed distance and fuel"],
        ["Time coverage",f"{s['paired_time_coverage']:.2%}","Valid paired / expected seconds"],
        ["Net fuel purchases",f"US${s['comparable_purchase_amount_usd']:,.2f}","Comparable transaction amount"],
    ],[1.65,1.8,3.75],family)
    figure(doc,root,fs[0],7.2)
    para(doc,"All fleet operations and follow-up cases are synthetic. Actual EIA observations are independent public background. These results demonstrate reproducible scenario analysis, not verified customer savings or production fleet performance.")
    newpage(doc,"Separate price effects from operating changes")
    figure(doc,root,fs[3],7.2)
    amounts=re.findall(r"US\$([\d,.]+)",fs[3]["findings"][0]);assert len(amounts)==5
    weeks=re.findall(r"\d{4}-\d{2}-\d{2}",fs[3]["findings"][0]);assert len(weeks)==2
    para(doc,f"Between the full weeks starting {weeks[0]} and {weeks[1]}, comparable purchases change from US${amounts[0]} to US${amounts[1]}. The price effect is US${amounts[2]}, the quantity effect is US${amounts[3]}, and the disclosed residual is US${amounts[4]}.")
    para(doc,"Fuel spending does not establish the cost of fuel consumed during the same week. Retrospective regional market gaps are reference comparisons; they are not executable offers or realizable savings.")
    para(doc,"Qualify comparisons before selecting cases","Heading 2")
    para(doc,f"The revised view identifies {s['low_coverage_rows']} vehicle-week below the coverage threshold and {s['no_eligible_peer_rows']} without an eligible peer baseline. Task-specific peer rules flag {s['intensity_review_rows']} vehicle-week; price rules flag {s['price_review_rows']}. Reasons may overlap.")
    para(doc,"Peers share tenant, vehicle class, task and week. Eligible comparisons need at least three other vehicles and 95% coverage. An intensity at least 20% above the eligible peer median calls for load, PTO and route review; it does not prove waste.")
    inv=statuses(fs[4]);known=dict(inv)
    para(doc,f"Inventory provides a second measurement path. Of {s['inventory_window_count']} windows, {known.get('within_measurement_uncertainty',0)} reconcile within stated measurement uncertainty. The remaining states preserve missing ECU intervals, uncalibrated boundaries, incomplete movements and cases needing coverage or balance investigation.")
    para(doc,"Use the offline dashboard for vehicle-week and task-level drilldown. Preserve unknown values and non-operating exposure rather than converting either into zero efficiency.")
    newpage(doc,"Follow adoption through comparable evidence")
    para(doc,f"The exercise includes {s['review_case_count']} review cases. Only {s['comparable_followup_cases']} pass the maturity, task and coverage checks for descriptive follow-up comparison. Rejected, unobserved, immature and lost cases retain their own states.")
    figure(doc,root,fs[6],7.2)
    cs=cases(fs[6]);assert len(cs)==s["comparable_followup_cases"]
    table(doc,["Vehicle","Baseline","Raw follow-up","Standardized","Change"],
          [[x[0],x[2],x[3],x[4],x[5]] for x in cs],[1.35,1.25,1.6,1.6,1.4],family)
    para(doc,"All rates are L/100km. Standardized follow-up rates use baseline task distance shares. Hollow diamonds show raw follow-up rates. Holding this measured task mix constant still does not control every operating difference.","Caption")
    para(doc,f"Recorded active review work totals {s['recorded_active_review_seconds']/60:.1f} minutes against a recorded baseline of {s['recorded_baseline_review_seconds']/60:.1f} minutes across {s['timed_review_cases']} synthetic cases. Waiting is excluded. This exercises the workload fields; it is not a measured productivity saving.")
    para(doc,"Reproduce and inspect","Heading 2")
    para(doc,"The repository contains source definitions, generation and acquisition code, SQL transformations, independent validation fixtures, an analytical database build, an executed notebook and the offline report. Follow README and docs/runbook.md to rebuild the full profile and change a reporting threshold.")
    para(doc,"Do not combine initial, revised and retrospective views. A 2026 retrieval of a 2025 price does not establish its historical availability. Synthetic outcomes and before/after changes must not be presented as real fleet effectiveness or causal savings.")
    path=root/"reports/DA-E01_Analysis_Report_EN.docx";doc.save(path);return path

def main():
    p=argparse.ArgumentParser();p.add_argument("--root",type=Path,default=ROOT);p.add_argument("--pdf",action="store_true");args=p.parse_args()
    root=args.root.resolve();(root/"reports").mkdir(exist_ok=True)
    source=root/"outputs/report_content.json";r=json.loads(source.read_text())
    paths=[zh_report(root,r),en_report(root,r)]
    if args.pdf:
        for path in paths:subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",str(root/"reports"),str(path)],check=True)
    print(json.dumps({"source":str(source),"documents":[str(p) for p in paths]},indent=2))

if __name__=="__main__":main()
